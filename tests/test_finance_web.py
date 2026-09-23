"""Раздел «Финансы» целиком: страницы, ввод, вложения, удаление, решения.

Владелец 23.09.2026: «сделать это всё серьёзно, чтобы не было каких-либо
багов». Поэтому здесь не рендер шаблона с выдуманным контекстом, а настоящий
кабинет на временной базе: запросы идут в приложение так же, как из браузера.
"""
import asyncio
import json
import os
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from urllib.parse import urlencode
from zoneinfo import ZoneInfo

os.environ.setdefault("OWNER_BOT_TOKEN", "test")
os.environ.setdefault("DRIVER_BOT_TOKEN", "test")
os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://u:p@localhost/db")
os.environ.setdefault("JWT_SECRET", "test-secret-please-set-a-long-one-in-prod")

import pytest  # noqa: E402

pytest.importorskip("aiosqlite")

from sqlalchemy import BigInteger, event, select  # noqa: E402
from sqlalchemy.dialects.postgresql import JSONB  # noqa: E402
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine  # noqa: E402
from sqlalchemy.ext.compiler import compiles  # noqa: E402

from app.models import (  # noqa: E402
    Base, Driver, Expense, ExpenseAttachment, ManualEntry, Owner, Shift, Trip, Vehicle,
)
from app.web import router as R  # noqa: E402


@compiles(JSONB, "sqlite")
def _jsonb(type_, compiler, **kw):
    return "JSON"


@compiles(BigInteger, "sqlite")
def _bigint(type_, compiler, **kw):
    return "INTEGER"


NOW = datetime.now(timezone.utc).replace(microsecond=0)
ZONE = ZoneInfo("Europe/Moscow")
TODAY = NOW.astimezone(ZONE).date()
MONTH_START = TODAY.replace(day=1)
PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
       b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\xff\xff?\x00\x05\xfe"
       b"\x02\xfe\xa7V\xbd\xfa\x00\x00\x00\x00IEND\xaeB`\x82")


# ── приложение поверх временной базы ─────────────────────────────────────────

class Cabinet:
    """Кабинет одного владельца: база своя, вход подставлен."""

    def __init__(self, maker, owner, data):
        self.maker = maker
        self.owner = owner
        self.data = data

    async def call(self, method: str, path: str, body: bytes = b"", ctype: str | None = None):
        path_only, _, query = path.partition("?")
        headers = [(b"host", b"testserver")]
        if ctype:
            headers.append((b"content-type", ctype.encode()))
            headers.append((b"content-length", str(len(body)).encode()))
        scope = {
            # spec 2.4: потоковый ответ (Excel) не ждёт «клиент ушёл» — иначе
            # наш receive() крутился бы в пустом цикле.
            "type": "http", "asgi": {"version": "3.0", "spec_version": "2.4"}, "http_version": "1.1",
            "method": method, "scheme": "http", "path": path_only, "raw_path": path_only.encode(),
            "query_string": query.encode(), "headers": headers,
            "client": ("127.0.0.1", 5000), "server": ("testserver", 80), "root_path": "",
            "app": R.app,
        }
        sent = {"done": False}
        out = {"status": None, "headers": {}, "body": b""}

        async def receive():
            # Тело — один раз; дальше «пусто», как у витрины (app/tools/showcase.py).
            if not sent["done"]:
                sent["done"] = True
                return {"type": "http.request", "body": body, "more_body": False}
            return {"type": "http.request", "body": b"", "more_body": False}

        async def send(message):
            if message["type"] == "http.response.start":
                out["status"] = message["status"]
                out["headers"] = {k.decode().lower(): v.decode() for k, v in message["headers"]}
            elif message["type"] == "http.response.body":
                out["body"] += message.get("body", b"")

        await R.app(scope, receive, send)
        return out

    async def get(self, path):
        return await self.call("GET", path)

    async def post_form(self, path, fields: dict):
        return await self.call("POST", path, urlencode(fields).encode(),
                               "application/x-www-form-urlencoded")

    async def post_multipart(self, path, fields: dict, files: list[tuple[str, str, str, bytes]] = ()):
        boundary = uuid.uuid4().hex
        parts = []
        for name, value in fields.items():
            parts.append(
                f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode()
            )
        for name, filename, ctype, data in files:
            parts.append(
                (f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"; '
                 f'filename="{filename}"\r\nContent-Type: {ctype}\r\n\r\n').encode() + data + b"\r\n"
            )
        parts.append(f"--{boundary}--\r\n".encode())
        return await self.call("POST", path, b"".join(parts), f"multipart/form-data; boundary={boundary}")


def _run(body):
    """Своя база и свой владелец на каждый тест."""
    async def scenario():
        engine = create_async_engine("sqlite+aiosqlite://", connect_args={"check_same_thread": False},
                                     poolclass=__import__("sqlalchemy.pool", fromlist=["StaticPool"]).StaticPool)

        @event.listens_for(engine.sync_engine, "connect")
        def _teach(conn, _):
            def tz(zone, moment):
                if moment is None:
                    return None
                when = datetime.fromisoformat(str(moment))
                if when.tzinfo is None:
                    when = when.replace(tzinfo=timezone.utc)
                return when.astimezone(ZoneInfo(str(zone))).strftime("%Y-%m-%d %H:%M:%S")
            conn.create_function("timezone", 2, tz)

        try:
            await _scenario_body(engine, body)
        finally:
            R.app.dependency_overrides.clear()
            await engine.dispose()

    asyncio.run(scenario())


async def _scenario_body(engine, body):
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        maker = async_sessionmaker(engine, expire_on_commit=False)

        async with maker() as s:
            owner = Owner(telegram_id=1, full_name="Владелец", company_name="ИП Кибиткина",
                          timezone="Europe/Moscow")
            stranger = Owner(telegram_id=2, full_name="Чужой", timezone="Europe/Moscow")
            s.add_all([owner, stranger])
            await s.flush()
            truck = Vehicle(owner_id=owner.id, license_plate="Т772НХ178", is_active=True)
            foreign_truck = Vehicle(owner_id=stranger.id, license_plate="А001АА78", is_active=True)
            s.add_all([truck, foreign_truck])
            await s.flush()
            driver = Driver(owner_id=owner.id, telegram_id=100, full_name="Саломов",
                            salary_type="per_km", salary_rate=Decimal(10))
            s.add(driver)
            await s.flush()
            shift = Shift(owner_id=owner.id, driver_id=driver.id, vehicle_id=truck.id,
                          status="started", started_at=NOW - timedelta(hours=5))
            s.add(shift)
            await s.flush()
            trip = Trip(owner_id=owner.id, driver_id=driver.id, vehicle_id=truck.id,
                        shift_id=shift.id, origin="Агропарк", destination="Магнит",
                        status="completed", completed_at=NOW - timedelta(hours=1),
                        revenue_rub=Decimal("19000"))
            s.add(trip)
            await s.flush()
            pending = Expense(owner_id=owner.id, driver_id=driver.id, shift_id=shift.id,
                              trip_id=trip.id, category="fuel", amount_rub=Decimal("7000"),
                              status="pending", created_at=NOW - timedelta(hours=2))
            approved = Expense(owner_id=owner.id, driver_id=driver.id, shift_id=shift.id,
                               category="toll", amount_rub=Decimal("640"), status="approved",
                               created_at=NOW - timedelta(hours=3))
            foreign_expense = Expense(owner_id=stranger.id, category="fuel",
                                      amount_rub=Decimal("999"), status="approved",
                                      created_by="owner", created_at=NOW)
            s.add_all([pending, approved, foreign_expense])
            await s.commit()
            data = {"truck": truck.id, "foreign_truck": foreign_truck.id, "driver": driver.id,
                    "shift": shift.id, "trip": trip.id, "pending": pending.id,
                    "approved": approved.id, "foreign_expense": foreign_expense.id,
                    "stranger": stranger.id}

        async def session_override():
            async with maker() as s:
                yield s

        async def owner_override():
            async with maker() as s:
                return await s.get(Owner, owner.id)

        R.app.dependency_overrides[R.get_session] = session_override
        R.app.dependency_overrides[R.current_owner] = owner_override
        await body(Cabinet(maker, owner, data))


def _json(resp):
    return json.loads(resp["body"].decode())


def _html(resp):
    assert resp["status"] == 200, (resp["status"], resp["body"][:400])
    return resp["body"].decode()


# ── страницы ────────────────────────────────────────────────────────────────

def test_overview_shows_money_charts_and_attention():
    async def body(c: Cabinet):
        html = _html(await c.get("/finances"))
        for text in ("Финансы", "Прибыль", "Денежный поток", "Куда уходят деньги",
                     "Топливо и остальное", "Машины", "Направления", "Последние операции"):
            assert text in html, text
        # Переключатель разделов и стрелка-раскрывашка в меню.
        assert 'aria-current="page">Обзор' in html
        assert 'popovertarget="navFinance"' in html and 'href="/finances/income"' in html
        # Трата водителя на решении видна сразу на обзоре.
        assert "1 трата ждёт решения" in html
        # Доход — выручка рейса; расход — только одобренное (640, не 7640).
        assert "19 000 ₽" in html
        assert "640 ₽" in html
        data = json.loads(html.split('<script id="fin-data" type="application/json">')[1].split("</script>")[0])
        assert sum(data["cashflow"]["income"]) == 19000
        assert sum(data["cashflow"]["expense"]) == 640
        assert data["categories"][0]["code"] == "toll"
    _run(body)


def test_expenses_page_lists_filters_and_trip_view():
    async def body(c: Cabinet):
        html = _html(await c.get("/finances/expenses"))
        assert "Сегодня" in html
        assert "ждёт решения" in html and 'data-decide="approve"' in html
        assert "Платная дорога" in html and "Топливо" in html
        assert 'id="sheet-expense"' in html and 'id="fin-item-tpl"' in html
        assert 'class="fin-seg__badge"' in html            # число трат на решении у вкладки

        only_fuel = _html(await c.get("/finances/expenses?cat=fuel"))
        assert "Платная дорога ·" not in only_fuel and "7 000" in only_fuel

        # Рейс прошлых месяцев не выглядит пустым: у рейса — все его траты.
        async with c.maker() as s:
            old = Expense(owner_id=c.owner.id, driver_id=c.data["driver"], trip_id=c.data["trip"],
                          category="parking", amount_rub=Decimal("300"), status="approved",
                          created_at=NOW - timedelta(days=80))
            s.add(old)
            await s.commit()
        trip_view = _html(await c.get(f"/finances/expenses?trip={c.data['trip']}"))
        assert "Парковка" in trip_view and "Агропарк → Магнит" in trip_view
        assert "Платная дорога" not in trip_view.split('id="sheet-expense"')[0]
    _run(body)


def test_income_page_has_trip_revenue_and_manual_entry():
    async def body(c: Cabinet):
        resp = await c.post_form("/finances/income", {
            "amount": "50 000", "entry_date": TODAY.isoformat(), "category": "Предоплата",
            "description": "за сентябрь",
        })
        assert resp["status"] == 200 and _json(resp)["ok"]
        html = _html(await c.get("/finances/income"))
        assert "Агропарк → Магнит" in html and "Предоплата" in html
        assert "69 000 ₽" in html                   # 19 000 рейс + 50 000 поступление
        bad = await c.post_form("/finances/income", {"amount": "0"})
        assert bad["status"] == 400 and "больше нуля" in _json(bad)["message"]
    _run(body)


# ── ввод трат ───────────────────────────────────────────────────────────────

def test_create_batch_with_photo_file_and_vehicle():
    async def body(c: Cabinet):
        items = [
            {"category": "fuel", "amount": "12 400,50", "description": "320 л"},
            {"category": "wash", "amount": "600", "description": ""},
        ]
        spent = (TODAY - timedelta(days=1)).isoformat() + "T09:30"
        resp = await c.post_multipart("/finances/expenses", {
            "items": json.dumps(items), "spent_at": spent, "link": "vehicle",
            "vehicle_id": str(c.data["truck"]), "payment_method": "fuel_card", "supplier": "Лукойл",
        }, files=[
            ("photos_0", "check.png", "image/png", PNG),
            ("files_1", "act.pdf", "application/pdf", b"%PDF-1.4 test"),
        ])
        assert resp["status"] == 200, resp["body"]
        ids = _json(resp)["ids"]
        assert len(ids) == 2
        async with c.maker() as s:
            rows = (await s.execute(select(Expense).where(Expense.id.in_(ids)).order_by(Expense.id))).scalars().all()
            assert [r.amount_rub for r in rows] == [Decimal("12400.50"), Decimal("600.00")]
            assert {r.vehicle_id for r in rows} == {c.data["truck"]}
            assert rows[0].batch_id and rows[0].batch_id == rows[1].batch_id
            assert rows[0].status == "approved" and rows[0].driver_id is None
            # Время — в поясе владельца: 09:30 по Москве = 06:30 UTC.
            assert rows[0].spent_at.replace(tzinfo=None) == datetime.fromisoformat(spent) - timedelta(hours=3)
            atts = (await s.execute(select(ExpenseAttachment).order_by(ExpenseAttachment.id))).scalars().all()
            assert [(a.expense_id, a.kind) for a in atts] == [(ids[0], "photo"), (ids[1], "file")]
            photo_id, file_id = atts[0].id, atts[1].id

        photo = await c.get(f"/finances/attachments/{photo_id}")
        assert photo["status"] == 200 and photo["body"] == PNG
        assert photo["headers"]["content-type"] == "image/png"
        assert photo["headers"]["content-disposition"].startswith("inline")
        doc = await c.get(f"/finances/attachments/{file_id}")
        assert doc["headers"]["x-content-type-options"] == "nosniff"

        # Новая трата видна в книге с отметкой «новая» и со скрепкой.
        html = _html(await c.get(f"/finances/expenses?period_from={TODAY - timedelta(days=3)}"
                                 f"&period_to={TODAY}&added={ids[0]},{ids[1]}"))
        assert html.count("fin-row is-new") == 2
        assert "Лукойл" in html and "Топливная карта" in html and "Мойка" in html
        assert f"/finances/attachments/{photo_id}" in html
        # «Вчера» — трата легла на день траты, а не на день внесения.
        assert "Вчера" in html
    _run(body)


def test_create_rejects_bad_input_with_human_messages():
    async def body(c: Cabinet):
        async def send(items, **extra):
            fields = {"items": json.dumps(items), "link": "company", **extra}
            resp = await c.post_multipart("/finances/expenses", fields)
            return resp["status"], _json(resp)["message"]

        assert await send([]) == (400, "Добавьте хотя бы одну трату.")
        assert (await send([{"category": "fuel", "amount": ""}]))[1] == "Трата №1: укажите сумму больше нуля."
        assert (await send([{"category": "", "amount": "5"}]))[1] == "Трата №1: выберите вид расхода."
        assert "«Прочего»" in (await send([{"category": "other", "amount": "5"}]))[1]
        assert "1 000 000" in (await send([{"category": "fuel", "amount": "1000001"}]))[1]
        status, text = await send([{"category": "fuel", "amount": "5"}], link="vehicle",
                                  vehicle_id=str(c.data["foreign_truck"]))
        assert (status, text) == (400, "Такой машины нет в парке.")
        status, text = await send([{"category": "fuel", "amount": "5"}],
                                  spent_at=(TODAY + timedelta(days=5)).isoformat() + "T10:00")
        assert status == 400 and "будущего" in text
        garbage = await c.post_multipart("/finances/expenses", {"items": "не json"})
        assert garbage["status"] == 400
        async with c.maker() as s:
            count = len((await s.execute(select(Expense).where(Expense.owner_id == c.owner.id))).all())
        assert count == 2                                   # ничего лишнего не записалось
    _run(body)


def test_expense_for_trip_inherits_shift_and_vehicle():
    async def body(c: Cabinet):
        resp = await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "parking", "amount": "350"}]),
            "link": "company", "trip_id": str(c.data["trip"]),
        })
        new_id = _json(resp)["ids"][0]
        async with c.maker() as s:
            e = await s.get(Expense, new_id)
            assert (e.trip_id, e.shift_id, e.vehicle_id) == (c.data["trip"], c.data["shift"], c.data["truck"])
        trip_page = _html(await c.get(f"/trips/{c.data['trip']}"))
        assert "Парковка" in trip_page and "внесён вами" in trip_page
        assert f"/finances/expenses?trip={c.data['trip']}&new=1" in trip_page
        shift_page = _html(await c.get(f"/shifts/{c.data['shift']}"))
        assert "Парковка" in shift_page and f"/finances/expenses?shift={c.data['shift']}&new=1" in shift_page
        # Лист ввода, открытый из рейса, уже привязан к нему.
        sheet = _html(await c.get(f"/finances/expenses?trip={c.data['trip']}&new=1"))
        assert 'data-open-new="expense"' in sheet
        assert f'name="trip_id" value="{c.data["trip"]}"' in sheet
    _run(body)


def test_custom_category_then_expense_with_it():
    async def body(c: Cabinet):
        resp = await c.post_form("/finances/categories", {"name": "  Тахограф  "})
        assert _json(resp) == {"ok": True, "code": "Тахограф", "label": "Тахограф"}
        again = await c.post_form("/finances/categories", {"name": "тахограф"})
        assert _json(again)["code"] == "Тахограф"
        empty = await c.post_form("/finances/categories", {"name": "   "})
        assert empty["status"] == 400
        resp = await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "Тахограф", "amount": "4500"}]), "link": "company",
        })
        assert resp["status"] == 200
        html = _html(await c.get("/finances/expenses"))
        assert "Тахограф" in html and 'value="Тахограф"' in html   # и в книге, и в листе ввода
    _run(body)


# ── удаление и решения ──────────────────────────────────────────────────────

def test_delete_own_expense_and_refuse_foreign():
    async def body(c: Cabinet):
        resp = await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "fuel", "amount": "100"}]), "link": "company",
        }, files=[("photos_0", "a.png", "image/png", PNG)])
        new_id = _json(resp)["ids"][0]
        assert (await c.call("POST", f"/finances/expenses/{new_id}/delete"))["status"] == 200
        assert (await c.call("POST", f"/finances/expenses/{c.data['foreign_expense']}/delete"))["status"] == 404
        async with c.maker() as s:
            assert await s.get(Expense, new_id) is None
            assert (await s.execute(select(ExpenseAttachment))).first() is None
            assert await s.get(Expense, c.data["foreign_expense"]) is not None
        # Чужое вложение по номеру не открывается.
        async with c.maker() as s:
            s.add(ExpenseAttachment(owner_id=c.data["stranger"], expense_id=c.data["foreign_expense"],
                                    kind="photo", data=PNG, size_bytes=len(PNG)))
            await s.commit()
            foreign_att = (await s.execute(select(ExpenseAttachment.id))).scalar_one()
        assert (await c.get(f"/finances/attachments/{foreign_att}"))["status"] == 404
    _run(body)


def test_decision_from_ledger_moves_money_into_totals():
    async def body(c: Cabinet):
        kpi = '<strong class="fin-kpi__value">{}\u00a0₽</strong>'
        before = _html(await c.get("/finances"))
        assert kpi.format("640") in before                  # расход — только одобренное
        resp = await c.post_form(f"/api/expenses/{c.data['pending']}/decision", {"action": "approve"})
        assert resp["status"] == 200 and _json(resp)["status"] == "approved"
        after = _html(await c.get("/finances"))
        assert kpi.format("7\u202f640") in after             # 640 + 7 000 одобренных
        assert "ждёт решения" not in after.split("Последние операции")[0]
    _run(body)


def test_manual_income_delete_via_old_route():
    async def body(c: Cabinet):
        entry_id = _json(await c.post_form("/finances/income", {"amount": "1000"}))["id"]
        assert (await c.call("POST", f"/finances/delete/{entry_id}"))["status"] == 200
        async with c.maker() as s:
            assert await s.get(ManualEntry, entry_id) is None
    _run(body)


# ── старые адреса и старые страницы ─────────────────────────────────────────

def test_old_addresses_lead_to_the_ledger():
    async def body(c: Cabinet):
        resp = await c.get("/expenses?category=fuel&status=pending&driver_id=3&date_from=2026-09-01")
        assert resp["status"] == 303
        loc = resp["headers"]["location"]
        assert loc.startswith("/finances/expenses?")
        for part in ("cat=fuel", "status=pending", "driver=3", "period_from=2026-09-01", "period_to="):
            assert part in loc
        assert (await c.get("/expenses"))["headers"]["location"] == "/finances/expenses"
        assert (await c.get("/fuel-history"))["headers"]["location"] == "/finances/expenses?cat=fuel"
    _run(body)


def test_edit_page_opens_for_owner_expense_and_keeps_kopecks():
    async def body(c: Cabinet):
        new_id = _json(await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "insurance", "amount": "1234,50"}]), "link": "company",
        }))["ids"][0]
        html = _html(await c.get(f"/expenses/{new_id}"))
        assert "Внесён вами" in html
        assert 'value="1234.50"' in html                     # раньше копейки срезались при сохранении
        assert '<option value="insurance" selected>Страховка</option>' in html
    _run(body)


def test_backdated_expense_counts_in_its_own_month():
    async def body(c: Cabinet):
        last_month_day = MONTH_START - timedelta(days=3)
        await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "repair", "amount": "25000"}]), "link": "company",
            "spent_at": last_month_day.isoformat() + "T12:00",
        })
        async with c.maker() as s:
            this_month = await R._finance_summary(s, c.owner.id, MONTH_START, TODAY, "Europe/Moscow")
            last_month = await R._finance_summary(s, c.owner.id, last_month_day.replace(day=1),
                                                  last_month_day, "Europe/Moscow")
        assert this_month["total_expense"] == Decimal("640")
        assert last_month["total_expense"] == Decimal("25000")
    _run(body)


def test_script_data_cannot_close_its_tag():
    async def body(c: Cabinet):
        async with c.maker() as s:
            s.add(Trip(owner_id=c.owner.id, shift_id=c.data["shift"], driver_id=c.data["driver"],
                       vehicle_id=c.data["truck"], origin="</script><b>", destination="X",
                       status="completed", completed_at=NOW - timedelta(minutes=5),
                       revenue_rub=Decimal("1")))
            await s.commit()
        html = _html(await c.get("/finances"))
        data_block = html.split('<script id="fin-data" type="application/json">')[1].split("</script>")[0]
        json.loads(data_block)                              # данные целые, тег не оборвался
        assert "&lt;/script&gt;&lt;b&gt; → X" in html          # в разметке — только как текст
    _run(body)


def test_excel_has_every_expense():
    async def body(c: Cabinet):
        import io
        from openpyxl import load_workbook

        await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "insurance", "amount": "41200", "description": "ОСАГО"}]),
            "link": "company", "supplier": "Ингосстрах", "payment_method": "transfer",
        })
        resp = await c.get(f"/finances/export.xlsx?period_from={MONTH_START}&period_to={TODAY}")
        assert resp["status"] == 200
        wb = load_workbook(io.BytesIO(resp["body"]))
        assert wb.sheetnames[:2] == ["Итог", "Расходы"]
        rows = [[c.value for c in row] for row in wb["Расходы"].iter_rows(min_row=2)]
        kinds = {r[1]: r for r in rows}
        assert kinds["Страховка"][2] == 41200 and kinds["Страховка"][5] == "вы"
        assert kinds["Страховка"][6] == "Безнал" and kinds["Страховка"][7] == "Ингосстрах"
        assert kinds["Топливо"][3] == "ждёт решения" and kinds["Топливо"][5] == "Саломов"
        summary = {r[0].value: r[1].value for r in wb["Итог"].iter_rows(min_row=2)}
        assert summary["Итого расход"] == 640 + 41200
    _run(body)


def test_trip_profit_follows_the_ledger():
    """Раньше траты рейса считались один раз — при завершении. Трата,
    внесённая потом кнопкой «+ Расход» в рейсе, отклонённая или удалённая,
    в прибыль рейса не попадала (найдено 23.09.2026)."""
    async def body(c: Cabinet):
        async def costs():
            async with c.maker() as s:
                t = await s.get(Trip, c.data["trip"])
                return t.fuel_cost_rub or 0, t.other_costs_rub or 0, t.profit_rub

        new_id = _json(await c.post_multipart("/finances/expenses", {
            "items": json.dumps([{"category": "parking", "amount": "350"}]),
            "link": "company", "trip_id": str(c.data["trip"]),
        }))["ids"][0]
        # Топливо водителя ждёт решения — по правилу рейса оно уже в расходах.
        assert await costs() == (Decimal("7000.00"), Decimal("350.00"), Decimal("11650.00"))

        await c.post_form(f"/api/expenses/{c.data['pending']}/decision", {"action": "reject"})
        assert await costs() == (Decimal("0.00"), Decimal("350.00"), Decimal("18650.00"))

        await c.call("POST", f"/finances/expenses/{new_id}/delete")
        assert await costs() == (Decimal("0.00"), Decimal("0.00"), Decimal("19000.00"))
    _run(body)
